#!/usr/bin/env python3
"""Generate PPE Detection presentation (PPTX) from project report content.

Run from repo root:
  pip install python-pptx
  python docs/generate_ppe_slides_pptx.py

Output: docs/ppe_detection_presentation.pptx
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN


def _add_bullet_slide(
    prs: Presentation,
    title: str,
    bullets: list[str | tuple[int, str]],
) -> None:
    """bullets: str for level 0, or (level, text) for indented bullets (level 0-2)."""
    layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1]
    tf = body.text_frame
    tf.clear()

    first = True
    for item in bullets:
        if isinstance(item, tuple):
            level, text = item
        else:
            level, text = 0, item
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.text = text
        p.level = min(level, 2)
        p.font.size = Pt(18)
        p.space_after = Pt(6)


def _add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    sub = slide.placeholders[1]
    sub.text = subtitle


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    _add_title_slide(
        prs,
        "PPE Detection with Fine-Tuned YOLOv8",
        "SH17 + APD merge • Head-only fine-tuning • Evaluation & API\n"
        "Python 3.11 • PyTorch 2.11 • Ultralytics 8.4.x • Apple MPS",
    )

    _add_bullet_slide(
        prs,
        "Project overview",
        [
            "Goal: detect personal protective equipment (PPE) end-to-end — data, training, metrics, API.",
            "Data: SH17 (17 classes) + APD v22 (Roboflow) → extended 18 classes (+ Face-shield).",
            "Preprocess: unify YOLO labels; resize longest side to 640; re-normalize every bbox.",
            "Model: YOLOv8m from COCO weights (yolov8m.pt).",
            "Training: --freeze 22 → train detection head only (backbone/neck frozen).",
            "Eval: inference.infer → mAP@0.5, mAP@0.5:0.95, P/R/F1, per-class AP50.",
            "Figures: inference.sample_predictions for val overlays (PPE-focused frames).",
        ],
    )

    _add_bullet_slide(
        prs,
        "Datasets: SH17 and APD",
        [
            "SH17: 17 PPE classes (CC BY 4.0), YOLO format; splits via train_files.txt / val_files.txt.",
            "Classes 0–16: Person, Head, Face, Glasses, Face-mask-medical, Face-guard, Ear, Earmuffs, Hands, Gloves, Foot, Shoes, Safety-vest, Tools, Helmet, Medical-suit, Safety-suit.",
            "APD v22: 9 original classes; includes face shield, mask, gloves, person, earmuff, and four “no-*” absence classes.",
            "APD chosen to add Face-shield annotations compatible with SH17-style bbox labels.",
            "Scale (planning doc): ~8,099 SH17 images; ~5,324 APD images — merged counts depend on filters when running prepare_data.py.",
        ],
    )

    _add_bullet_slide(
        prs,
        "Merging SH17 + APD (prepare_data.py)",
        [
            "Extended: SH17 train+val + APD train, valid, and APD test images merged into val.",
            "Prefixes avoid collisions: sh17_<stem>, apd_<stem>.",
            "SH17 labels: exactly 5 fields per line; class ∈ [0,16]; coords clamped to [0,1].",
            "APD remap: 0→Earmuffs(7), 1→Face-shield(17), 2→Gloves(9), 3→Face-mask-medical(4), 8→Person(0).",
            "Discard APD negative classes {4,5,6,7} (“no earmuff/shield/glove/mask”).",
            "Polygon lines (>4 coords): axis-aligned bbox via min/max of vertices (~12.4% of APD lines).",
            "Outputs: data.yaml + merge_report.json (counts, skips, class distribution).",
            "Modes: --sh17-only (17-class baseline dir); --sample (50 SH17 + 20 APD per split, ordered).",
        ],
    )

    _add_bullet_slide(
        prs,
        "Resize pipeline (resize_dataset.py)",
        [
            "Proportional resize: max(width, height) = max_dim (default 640); preserve aspect ratio.",
            "If already small enough: scale 1.0; else scale = max_dim / max(W_orig, H_orig).",
            "W_new = round(W_orig·s), H_new = round(H_orig·s); PIL LANCZOS; JPEG quality default 95; RGB.",
            "BBox math per box: denorm to orig pixels → multiply by s → re-normalize to W_new, H_new.",
            "Drop boxes with w or h < 1e-4 after resize (tiny box filter).",
            "Training YAMLs: sh17_baseline_resized.yaml → data/sh17_baseline_resized; sh17_extended_resized.yaml → data/merged_resized.",
        ],
    )

    _add_bullet_slide(
        prs,
        "YOLOv8m architecture & freezing",
        [
            "Backbone: CSP-style feature extractor; Neck: FPN/PAN fusion; Head: decoupled, anchor-free, 3 scales (P3/P4/P5).",
            "At imgsz=640: stride 8/16/32 → grids ~80×80, 40×40, 20×20.",
            "Input tensor shape: B × 3 × 640 × 640.",
            "Each scale predicts dense box distributions + class logits; decode + NMS → final boxes.",
            "This project: model.train(..., freeze=22) — first 22 layers frozen; only head (tail) trains.",
            "Rationale: reuse COCO visual features; adapt class logits and regression for PPE taxonomy quickly.",
        ],
    )

    _add_bullet_slide(
        prs,
        "How YOLO trains here (losses & recipe)",
        [
            "Box: IoU-style loss (e.g. CIoU) on decoded boxes + Distribution Focal Loss (DFL) on edge distributions.",
            "Classification: BCE-with-logits with task-aligned positive/negative assignment.",
            "TaskAlignedAssigner (TAL): matches preds to GT; repo patches TAL for MPS stability (training/tal_patch.py).",
            "Train via: python training/train.py from repo root (do not skip tal_patch).",
            "Hyperparameters (documented runs): yolov8m.pt, lr0=0.01, cos LR, batch 32, imgsz=640, workers 4, --no-amp on MPS.",
            "mosaic=0, close_mosaic=0; val every 5 epochs; max_det=100 for validation speed.",
        ],
    )

    _add_bullet_slide(
        prs,
        "Checkpoints: pretrained vs baseline_20ep vs extended_20ep",
        [
            "Pretrained: yolov8m.pt (COCO, 80 classes) — before PPE adaptation.",
            "Note: evaluating pretrained with SH17 YAML does not yield meaningful “PPE” per-class names; use as before/after story.",
            "baseline_20ep: best.pt on sh17_baseline_resized (17 classes).",
            "extended_20ep: best.pt on sh17_extended_resized (18 classes, + Face-shield).",
            "Same training recipe; only dataset YAML / nc differs.",
        ],
    )

    _add_bullet_slide(
        prs,
        "Validation metrics (val, imgsz 640, conf 0.25, IoU 0.7)",
        [
            "Pretrained (yolov8m.pt on SH17 YAML): mAP@0.5 0.054, mAP@0.5:0.95 0.048, F1 ~0.052.",
            "baseline_20ep: mAP@0.5 0.505, mAP@0.5:0.95 0.322, P 0.556, R 0.421, F1 0.458.",
            "extended_20ep: mAP@0.5 0.516, mAP@0.5:0.95 0.326, P 0.568, R 0.423, F1 0.457.",
            "Face-shield (extended): AP@0.5 ≈ 0.40 (see results/extended_20ep/metrics_finetuned.json).",
            "Shared-class example deltas (AP50): Person 0.82→0.79, Face 0.20→0.19, Foot 0.36→0.29.",
            "Latency (ms/img, JSON speed): pretrained ~0.22/19.9/13.8 preprocess/infer/post; extended_20ep ~0.21/20.6/14.3.",
            "Figures: bar chart of per-class AP50 (17 shared + Face-shield); diagram B×3×640×640 → 80/40/20 grids.",
        ],
    )

    _add_bullet_slide(
        prs,
        "Qualitative: images through YOLO",
        [
            "Script: python -m inference.sample_predictions --model <weights> --data <yaml> --project … --name …",
            "Selects val images with SH17 equipment ids 1–16 and/or Face-shield (17) on extended data.",
            "Compare side-by-side: pretrained (COCO labels) vs baseline vs extended on the same frame.",
            "Suggested output dirs: results/baseline/samples/*, results/extended/samples/*.",
        ],
    )

    _add_bullet_slide(
        prs,
        "API, reproducibility, wrap-up",
        [
            "FastAPI: POST /auth/token (JWT), /predict/image, /predict/batch; GET /health, /metrics.",
            "Docker: docker compose up --build (see README).",
            "Evaluation: python -m inference.infer --model … --data … --output …",
            "Rubric compare: python -m inference.compare (baseline vs extended JSON).",
            "Always train with python training/train.py so TAL patch applies on MPS.",
        ],
    )

    return prs


def main() -> None:
    out = Path(__file__).resolve().parent / "ppe_detection_presentation.pptx"
    prs = build_presentation()
    prs.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
